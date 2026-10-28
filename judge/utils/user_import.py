"""Bulk account import from an Excel/CSV file, shared by the bulk user pages and the contestant import popup.

Flow: ``read_file`` -> ``resolve_rows`` (match existing accounts by email, plan usernames) -> the caller sets each
row's ``status`` -> ``create_accounts`` for the rows marked ``STATUS_CREATE``.
"""
import csv
import io
import logging
import re
import secrets
import unicodedata
import uuid
import zipfile
from dataclasses import dataclass

from django.conf import settings
from django.contrib.auth.models import User
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from django.db.models.functions import Lower
from django.urls import reverse
from django.utils.translation import gettext as _

from judge.models import Language, Profile

logger = logging.getLogger('judge.user_import')

MAX_ROWS = 500
COLUMNS = ('email', 'username', 'fullname', 'password')
USERNAME_RE = re.compile(r'^\w+$', re.ASCII)
USERNAME_MAX_LENGTH = 30
FULLNAME_MAX_LENGTH = User._meta.get_field('first_name').max_length
PASSWORD_LENGTH = 10
# No 0/O, 1/l/I: generated passwords are often typed from a printout.
PASSWORD_ALPHABET = 'abcdefghjkmnpqrstuvwxyz' 'ABCDEFGHJKLMNPQRSTUVWXYZ' '23456789'

STATUS_CREATE = 'create'  # a new account will be created
STATUS_ADD = 'add'  # an existing account will be added (to a contest or organization)
STATUS_SKIP = 'skip'  # nothing to do for this row
STATUS_ERROR = 'error'

CACHE_TIMEOUT = 60 * 60


class ImportFileError(Exception):
    pass


@dataclass
class ImportRow:
    line: int
    email: str
    username: str = ''
    fullname: str = ''
    password: str = ''
    error: str = ''
    status: str = ''
    note: str = ''
    # The account the email belongs to, or the one created by create_accounts.
    profile: Profile = None

    @property
    def is_new(self):
        return self.profile is None and not self.error

    def to_cache(self):
        return {key: getattr(self, key) for key in ('line',) + COLUMNS}


def _cell_text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        # Excel stores numeric passwords like 123456 as 123456.0.
        value = int(value)
    return str(value).strip()


def _read_xlsx(uploaded_file):
    import openpyxl

    try:
        workbook = openpyxl.load_workbook(uploaded_file, read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError):
        raise ImportFileError(_('Could not read the Excel file.'))
    try:
        return [[_cell_text(cell) for cell in row] for row in workbook.worksheets[0].iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_csv(uploaded_file):
    try:
        text = uploaded_file.read().decode('utf-8-sig')
    except UnicodeDecodeError:
        raise ImportFileError(_('The CSV file must be UTF-8 encoded.'))
    return [[_cell_text(cell) for cell in row] for row in csv.reader(io.StringIO(text))]


def read_file(uploaded_file):
    """Read the rows of an uploaded .xlsx or .csv file whose first row has the column names."""
    name = (uploaded_file.name or '').lower()
    if name.endswith('.xlsx'):
        table = _read_xlsx(uploaded_file)
    elif name.endswith('.csv'):
        table = _read_csv(uploaded_file)
    else:
        raise ImportFileError(_('Only .xlsx and .csv files are supported.'))

    if not table:
        raise ImportFileError(_('The file is empty.'))
    header = [cell.lower() for cell in table[0]]
    if 'email' not in header:
        raise ImportFileError(_('The file has no "email" column. The first row must contain the column names: %s.')
                              % ', '.join(COLUMNS))
    index = {column: header.index(column) for column in COLUMNS if column in header}

    rows = []
    for line, cells in enumerate(table[1:], start=2):
        if not any(cells):
            continue
        values = {column: cells[i] if i < len(cells) else '' for column, i in index.items()}
        rows.append(ImportRow(line=line, **values))
    if not rows:
        raise ImportFileError(_('The file has no rows besides the column names.'))
    if len(rows) > MAX_ROWS:
        raise ImportFileError(_('The file has %(count)d rows; at most %(max)d are allowed at once.')
                              % {'count': len(rows), 'max': MAX_ROWS})
    return rows


def generate_password():
    return ''.join(secrets.choice(PASSWORD_ALPHABET) for _ in range(PASSWORD_LENGTH))


def username_base(email):
    local = email.split('@')[0].replace('đ', 'd').replace('Đ', 'D')
    local = ''.join(c for c in unicodedata.normalize('NFKD', local) if not unicodedata.combining(c)).lower()
    return (re.sub(r'[^a-z0-9_]', '', local) or 'user')[:USERNAME_MAX_LENGTH - 4]


def _generate_username(email, reserved):
    base = username_base(email)
    taken = set(User.objects.annotate(name=Lower('username')).filter(name__startswith=base)
                .values_list('name', flat=True)) | reserved
    username, suffix = base, 1
    while username in taken:
        suffix += 1
        username = '%s%d' % (base, suffix)
    return username


def resolve_rows(rows):
    """Validate the rows, attach the existing account of each email and plan usernames for new accounts."""
    emails = {row.email.lower() for row in rows if row.email}
    accounts = {}
    for profile in Profile.objects.select_related('user').annotate(email_lower=Lower('user__email')) \
            .filter(email_lower__in=emails):
        accounts.setdefault(profile.email_lower, []).append(profile)

    wanted = {row.username.lower() for row in rows if row.username}
    taken = set(User.objects.annotate(name=Lower('username')).filter(name__in=wanted)
                .values_list('name', flat=True))

    seen, reserved = {}, set()
    for row in rows:
        row.error, row.profile = '', None
        if not row.email:
            row.error = _('Missing email.')
            continue
        try:
            validate_email(row.email)
        except ValidationError:
            row.error = _('Invalid email address.')
            continue
        key = row.email.lower()
        if key in seen:
            row.error = _('Duplicate of the email on line %d.') % seen[key]
            continue
        seen[key] = row.line

        matches = accounts.get(key, [])
        if len(matches) > 1:
            row.error = _('Several accounts use this email.')
        elif matches:
            row.profile = matches[0]
        elif len(row.fullname) > FULLNAME_MAX_LENGTH:
            row.error = _('Full name is longer than %d characters.') % FULLNAME_MAX_LENGTH
        elif row.username:
            name = row.username.lower()
            if not USERNAME_RE.match(row.username) or len(row.username) > USERNAME_MAX_LENGTH:
                row.error = _('A username may only contain letters, numbers and underscores (at most %d).') \
                    % USERNAME_MAX_LENGTH
            elif name in taken or name in reserved:
                row.error = _('Username %s is already taken.') % row.username
            else:
                reserved.add(name)
        else:
            row.username = _generate_username(row.email, reserved)
            reserved.add(row.username)
    return rows


def create_accounts(rows, organization=None):
    """Create the accounts of the rows marked STATUS_CREATE; their ``password`` is set to the one used."""
    language = Language.get_default_language()
    created = []
    with transaction.atomic():
        for row in rows:
            if row.status != STATUS_CREATE:
                continue
            row.password = row.password or generate_password()
            user = User(username=row.username, email=row.email, first_name=row.fullname, is_active=True)
            user.set_password(row.password)
            user.save()
            row.profile = Profile.objects.create(user=user, language=language)
            if organization is not None:
                row.profile.organizations.add(organization)
            created.append(row)
        if organization is not None and created:
            organization.on_user_changes()
    return created


def credentials_workbook(rows):
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'accounts'
    sheet.append(['username', 'fullname', 'email', 'password'])
    for row in rows:
        sheet.append([row.username, row.fullname, row.email, row.password])
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def template_workbook():
    """An empty file to fill in: the column names on the first sheet and what they mean on the second."""
    import openpyxl
    from openpyxl.styles import Font

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = 'users'
    sheet.append(list(COLUMNS))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for letter, width in zip('ABCD', (32, 20, 28, 16)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = 'A2'

    guide = workbook.create_sheet(_('guide')[:31])
    guide.append([_('Column'), _('Description')])
    for cell in guide[1]:
        cell.font = Font(bold=True)
    guide.append(['email', _('Required. Rows without a valid email are reported as errors.')])
    guide.append(['username', _('Optional. Generated from the email when empty.')])
    guide.append(['fullname', _('Optional.')])
    guide.append(['password', _('Optional. A random password is generated when empty.')])
    guide.append([])
    guide.append([_('If an email already has an account, that account is used and the username, fullname and '
                    'password in the file are ignored.')])
    guide.append([_('Fill in the "users" sheet from its second row; only that sheet is read.')])
    guide.column_dimensions['A'].width = 14
    guide.column_dimensions['B'].width = 70

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def send_credentials_emails(rows, request):
    """Email each new account its login details. Returns the usernames that could not be emailed."""
    failed = []
    login_url = request.build_absolute_uri(reverse('auth_login'))
    for row in rows:
        subject = _('Your account on %s') % settings.SITE_NAME
        message = _('Hello %(name)s,\n\n'
                    'An account has been created for you on %(site)s.\n\n'
                    'Username: %(username)s\n'
                    'Password: %(password)s\n\n'
                    'Log in at %(url)s and change your password afterwards.') % {
            'name': row.fullname or row.username, 'site': settings.SITE_LONG_NAME,
            'username': row.username, 'password': row.password, 'url': login_url,
        }
        try:
            row.profile.user.email_user(subject, message)
        except Exception:
            logger.exception('Failed to email the new account %s', row.username)
            failed.append(row.username)
    return failed


def _cache_key(kind, user, token):
    return 'user_import_%s:%d:%s' % (kind, user.id, token)


def store_rows(user, rows):
    """Keep the uploaded rows between the preview and the confirmation; returns the token to confirm with."""
    token = uuid.uuid4().hex
    cache.set(_cache_key('rows', user, token), [row.to_cache() for row in rows], CACHE_TIMEOUT)
    return token


def load_rows(user, token):
    data = cache.get(_cache_key('rows', user, token)) if token else None
    if data is None:
        return None
    cache.delete(_cache_key('rows', user, token))
    return [ImportRow(**item) for item in data]


def store_credentials(user, rows):
    token = uuid.uuid4().hex
    cache.set(_cache_key('credentials', user, token), credentials_workbook(rows), CACHE_TIMEOUT)
    return token


def load_credentials(user, token):
    return cache.get(_cache_key('credentials', user, token))


def row_summary(rows):
    counts = {STATUS_CREATE: 0, STATUS_ADD: 0, STATUS_SKIP: 0, STATUS_ERROR: 0}
    for row in rows:
        counts[row.status] += 1
    return counts
