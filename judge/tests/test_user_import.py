import io

import openpyxl
from django.contrib.auth.models import User
from django.core import mail
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from judge.models import Profile
from judge.models.tests.util import create_contest, create_organization, create_user
from judge.utils.user_import import COLUMNS, ImportFileError, ImportRow, MAX_ROWS, STATUS_CREATE, \
    create_accounts, read_file, resolve_rows, username_base


def xlsx_file(rows, name='users.xlsx'):
    workbook = openpyxl.Workbook()
    for row in rows:
        workbook.active.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return SimpleUploadedFile(name, output.getvalue())


def csv_file(text, name='users.csv'):
    return SimpleUploadedFile(name, text.encode('utf-8-sig'))


class ReadFileTestCase(TestCase):
    def test_reads_xlsx_with_any_column_order(self):
        rows = read_file(xlsx_file([
            ['fullname', 'Email', 'password'],
            ['Nguyễn Văn A', 'a@example.com', 123456],
            [None, None, None],
            ['B', 'b@example.com', None],
        ]))
        self.assertEqual([(r.line, r.email, r.fullname, r.password) for r in rows],
                         [(2, 'a@example.com', 'Nguyễn Văn A', '123456'), (4, 'b@example.com', 'B', '')])

    def test_reads_csv(self):
        rows = read_file(csv_file('email,username\na@example.com,alice\n'))
        self.assertEqual((rows[0].email, rows[0].username), ('a@example.com', 'alice'))

    def test_rejects_bad_files(self):
        for upload in (
            csv_file('username,fullname\nalice,Alice\n'),  # no email column
            csv_file('email\n'),  # no rows
            SimpleUploadedFile('users.txt', b'email\na@example.com'),
            SimpleUploadedFile('users.xlsx', b'not a zip file'),
            csv_file('email\n' + '\n'.join('u%d@example.com' % i for i in range(MAX_ROWS + 1))),
        ):
            with self.subTest(name=upload.name), self.assertRaises(ImportFileError):
                read_file(upload)


class TemplateTestCase(TestCase):
    def test_template_round_trip(self):
        self.client.force_login(create_user(username='someone'))
        response = self.client.get(reverse('user_import_template'))
        self.assertIn('attachment', response['Content-Disposition'])

        workbook = openpyxl.load_workbook(io.BytesIO(response.content))
        self.assertEqual(next(workbook.worksheets[0].iter_rows(values_only=True)), COLUMNS)
        self.assertEqual(len(workbook.worksheets), 2)

        # Filled in, the template reads back; the guide sheet is ignored.
        workbook.worksheets[0].append(['a@example.com', None, 'A', None])
        output = io.BytesIO()
        workbook.save(output)
        rows = read_file(SimpleUploadedFile('filled.xlsx', output.getvalue()))
        self.assertEqual([(r.email, r.fullname) for r in rows], [('a@example.com', 'A')])


class ResolveRowsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.existing = create_user(username='existing', email='Existing@Example.com')
        create_user(username='nguyenvana')

    def resolve(self, *rows):
        return resolve_rows([ImportRow(line=i + 2, **row) for i, row in enumerate(rows)])

    def test_username_base(self):
        self.assertEqual(username_base('Nguyễn.Văn-Đức@example.com'), 'nguyenvanduc')
        self.assertEqual(username_base('...@example.com'), 'user')

    def test_rows(self):
        rows = self.resolve(
            {'email': ''},
            {'email': 'not-an-email'},
            {'email': 'existing@example.com', 'username': 'ignored'},
            {'email': 'nguyen.van.a@school.edu'},
            {'email': 'nguyen.van.a@other.edu'},
            {'email': 'NGUYEN.VAN.A@school.edu'},
            {'email': 'c@example.com', 'username': 'existing'},
            {'email': 'd@example.com', 'username': 'bad name'},
            {'email': 'e@example.com', 'username': 'chosen'},
        )
        self.assertTrue(rows[0].error and rows[1].error)
        self.assertEqual(rows[2].profile.user, self.existing)
        self.assertEqual((rows[3].username, rows[4].username), ('nguyenvana2', 'nguyenvana3'))
        self.assertIn('line 5', rows[5].error)
        self.assertIn('already taken', rows[6].error)
        self.assertTrue(rows[7].error)
        self.assertEqual((rows[8].username, rows[8].error), ('chosen', ''))

    def test_create_accounts(self):
        org = create_organization(name='school')
        rows = self.resolve({'email': 'new@example.com', 'fullname': 'New', 'password': 'secret123'},
                            {'email': 'gen@example.com'})
        for row in rows:
            row.status = STATUS_CREATE
        create_accounts(rows, organization=org)

        user = User.objects.get(username='new')
        self.assertTrue(user.check_password('secret123'))
        self.assertEqual((user.email, user.first_name), ('new@example.com', 'New'))
        self.assertTrue(user.profile.is_approved)
        self.assertTrue(user.profile.must_change_password)
        self.assertEqual(len(rows[1].password), 10)
        self.assertTrue(User.objects.get(username='gen').check_password(rows[1].password))
        self.assertEqual(set(org.members.values_list('user__username', flat=True)), {'new', 'gen'})


class ImportViewTestMixin:
    def upload(self, url, upload, current=None):
        data = {'step': 'preview', 'file': upload}
        if current is not None:
            data['current'] = current
        return self.client.post(url, data)

    def confirm(self, url, preview, send_email=False, current=None):
        data = {'step': 'confirm', 'token': preview.context['token']}
        if send_email:
            data['send_email'] = '1'
        if current is not None:
            data['current'] = current
        return self.client.post(url, data)

    def statuses(self, response):
        return {row.email: row.status for row in response.context['rows']}


class BulkUserImportTestCase(ImportViewTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.importer = create_user(username='importer', user_permissions=('bulk_create_users',))
        cls.regular = create_user(username='regular')
        create_user(username='old', email='old@example.com')

    def test_needs_permission(self):
        self.client.force_login(self.regular)
        self.assertEqual(self.client.get(reverse('user_import')).status_code, 403)

    def test_preview_confirm_download_and_email(self):
        self.client.force_login(self.importer)
        url = reverse('user_import')
        self.assertContains(self.client.get(url), reverse('user_import_template'))
        preview = self.upload(url, csv_file('email,fullname\nnew@example.com,New\nold@example.com,Old\n,x\n'))
        self.assertEqual(self.statuses(preview), {'new@example.com': 'create', 'old@example.com': 'skip', '': 'error'})
        self.assertFalse(User.objects.filter(email='new@example.com').exists())

        result = self.confirm(url, preview, send_email=True)
        self.assertTrue(User.objects.filter(username='new', email='new@example.com').exists())
        self.assertEqual([m.to for m in mail.outbox], [['new@example.com']])
        password = result.context['created'][0].password
        self.assertIn(password, mail.outbox[0].body)

        download = self.client.get(result.context['credentials_url'])
        sheet = openpyxl.load_workbook(io.BytesIO(download.content)).active
        self.assertEqual(list(sheet.iter_rows(values_only=True)),
                         [('username', 'fullname', 'email', 'password'), ('new', 'New', 'new@example.com', password)])

        # The preview token is single-use.
        self.assertEqual(self.confirm(url, preview).context['step'], 'upload')


class OrganizationUserImportTestCase(ImportViewTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.admin = create_user(username='orgadmin')
        cls.org = create_organization(name='school', admins=('orgadmin',))
        cls.superuser = create_user(username='root', is_superuser=True)
        cls.outsider = create_user(username='outsider', email='outsider@example.com')
        cls.url = reverse('organization_user_import', args=['school'])

    def test_org_admin_creates_members_but_leaves_existing_accounts(self):
        self.client.force_login(self.admin)
        preview = self.upload(self.url, csv_file('email\nkid@example.com\noutsider@example.com\n'))
        self.assertEqual(self.statuses(preview), {'kid@example.com': 'create', 'outsider@example.com': 'skip'})
        self.confirm(self.url, preview)
        # The admin was already a member; outsider wasn't pulled in.
        self.assertEqual(set(self.org.members.values_list('user__username', flat=True)), {'orgadmin', 'kid'})

    def test_superuser_also_adds_existing_accounts(self):
        self.client.force_login(self.superuser)
        preview = self.upload(self.url, csv_file('email\noutsider@example.com\n'))
        self.assertEqual(self.statuses(preview), {'outsider@example.com': 'add'})
        self.confirm(self.url, preview)
        self.assertIn(self.org, self.outsider.profile.organizations.all())

    def test_needs_org_admin(self):
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(self.url).status_code, 403)


class ContestantImportTestCase(ImportViewTestMixin, TestCase):
    @classmethod
    def setUpTestData(cls):
        # Like the contest edit page, editing needs edit_own_contest on top of being an author.
        cls.author = create_user(username='author', user_permissions=('edit_own_contest',))
        cls.creator = create_user(username='creator', user_permissions=('edit_own_contest', 'bulk_create_users'))
        cls.listed = create_user(username='listed', email='listed@example.com')
        cls.other = create_user(username='other', email='other@example.com')
        cls.contest = create_contest(key='open', authors=('author', 'creator'))
        cls.url = reverse('contest_contestant_import', args=['open'])

        cls.org_admin = create_user(username='orgadmin', user_permissions=('edit_own_contest',))
        cls.org = create_organization(name='school', admins=('orgadmin',))
        cls.member = create_user(username='member', email='member@example.com')
        cls.member.profile.organizations.add(cls.org)
        cls.org_contest = create_contest(key='school_c', authors=('orgadmin',), is_organization_private=True,
                                         organization=cls.org)
        cls.org_url = reverse('contest_contestant_import', args=['school_c'])

    def test_editor_without_bulk_permission_adds_existing_only(self):
        self.client.force_login(self.author)
        current = str(self.listed.profile.id)
        preview = self.upload(self.url, csv_file('email\nlisted@example.com\nother@example.com\nnew@example.com\n'),
                              current=current)
        self.assertEqual(self.statuses(preview),
                         {'listed@example.com': 'skip', 'other@example.com': 'add', 'new@example.com': 'error'})

        result = self.confirm(self.url, preview, current=current)
        self.assertEqual(result.context['contestants'], [{'id': self.other.profile.id, 'text': 'other'}])
        self.assertFalse(User.objects.filter(email='new@example.com').exists())
        # The contest itself is only saved by the contest form.
        self.assertFalse(self.contest.private_contestants.exists())

    def test_creator_creates_accounts(self):
        self.client.force_login(self.creator)
        preview = self.upload(self.url, csv_file('email\nnew@example.com\n'))
        result = self.confirm(self.url, preview)
        new = Profile.objects.get(user__email='new@example.com')
        self.assertEqual(result.context['contestants'], [{'id': new.id, 'text': 'new'}])

    def test_org_contest_requires_membership(self):
        self.client.force_login(self.org_admin)
        preview = self.upload(self.org_url,
                              csv_file('email\nmember@example.com\nother@example.com\nkid@example.com\n'))
        self.assertEqual(self.statuses(preview),
                         {'member@example.com': 'add', 'other@example.com': 'error', 'kid@example.com': 'create'})
        result = self.confirm(self.org_url, preview)
        kid = Profile.objects.get(user__email='kid@example.com')
        self.assertIn(self.org, kid.organizations.all())
        self.assertEqual({c['text'] for c in result.context['contestants']}, {'member', 'kid'})

    def test_popup_can_be_framed_by_the_site(self):
        self.client.force_login(self.author)
        self.assertEqual(self.client.get(self.url)['X-Frame-Options'], 'SAMEORIGIN')

    def test_needs_edit_permission(self):
        self.client.force_login(self.other)
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.force_login(self.author)
        self.assertEqual(self.client.get(reverse('organization_contestant_import', args=['school'])).status_code,
                         403)


class BatchAddUserCommandTestCase(TestCase):
    def test_creates_accounts_and_writes_passwords(self):
        import csv
        import os
        import tempfile
        from django.core.management import call_command

        create_user(username='old', email='old@example.com')
        with tempfile.TemporaryDirectory() as tmp:
            source, output = os.path.join(tmp, 'in.csv'), os.path.join(tmp, 'out.csv')
            with open(source, 'w') as f:
                f.write('email,fullname\nnew@example.com,New\nold@example.com,Old\n')
            call_command('batchadduser', source, output, stdout=io.StringIO(), stderr=io.StringIO())
            with open(output) as f:
                written = list(csv.DictReader(f))
        self.assertEqual([(r['username'], r['email']) for r in written], [('new', 'new@example.com')])
        self.assertTrue(User.objects.get(username='new').check_password(written[0]['password']))


@override_settings(AUTH_PASSWORD_VALIDATORS=[])
class MustChangePasswordTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.user = create_user(username='student')
        cls.user.set_password('generated1')
        cls.user.save()
        Profile.objects.filter(user=cls.user).update(must_change_password=True)

    def test_redirects_until_the_password_is_changed(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse('problem_list'))
        self.assertRedirects(response, reverse('password_change') + '?next=' + reverse('problem_list'),
                             fetch_redirect_response=False)
        self.assertEqual(self.client.get(reverse('password_change')).status_code, 200)

        self.client.post(reverse('password_change'), {
            'old_password': 'generated1', 'new_password1': 'my-own-pass-2', 'new_password2': 'my-own-pass-2',
        })
        self.assertFalse(Profile.objects.get(user=self.user).must_change_password)
        self.assertEqual(self.client.get(reverse('problem_list')).status_code, 200)

    def test_any_new_password_clears_the_flag(self):
        user = User.objects.get(id=self.user.id)
        user.last_login = None
        user.save(update_fields=['last_login'])
        self.assertTrue(Profile.objects.get(user=user).must_change_password)

        user.set_password('reset-by-email-1')  # e.g. the password reset link
        user.save()
        self.assertFalse(Profile.objects.get(user=user).must_change_password)

    @override_settings(VNOJ_OFFICIAL_CONTEST_MODE=True)
    def test_not_enforced_in_official_contest_mode(self):
        self.client.force_login(self.user)
        self.assertEqual(self.client.get(reverse('problem_list')).status_code, 200)
